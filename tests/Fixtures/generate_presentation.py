"""Regenerate the one-slide Office smoke-test fixture (requires python-pptx)."""
from pathlib import Path
from pptx import Presentation

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[0])
slide.shapes.title.text = "Orbit PDF export"
slide.placeholders[1].text = "One-slide conversion fixture"
presentation.save(Path(__file__).with_name("presentation.pptx"))
